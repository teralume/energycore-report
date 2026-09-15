"""Build the AV1 participation report using the exact Annex B table structure."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "upc-pre-202610-1asi0732-9100-teralume-performance-av1.docx"

# Names, codes and institutional emails were provided by the team. The team
# leader confirmed that every listed member supported AV1 and requested on-time
# marks and scores of 20. Attribute only tasks evidenced for each individual.
MEMBERS = [
    {
        "student": "Loa Rojas, Jean Franck\nU20241E406",
        "responsibilities": [
            "Integración de los repositorios, arquitectura de la API y adaptación del producto EnergyCore.",
            "Desarrollo y verificación de la aplicación web, Landing Page y aplicación móvil Flutter.",
            "Documentación del proyecto, Student Outcome ABET y evidencias técnicas de AV1.",
        ],
        "score": "20",
    },
    {
        "student": "Jairo Mathias Santiago Atanacio\nU202418755\nu202418755@upc.edu.pe",
        "responsibilities": [
            "Incorporación y actualización de su perfil en el Project Report.",
            "Carga de su fotografía y corrección del formato y de una entrada duplicada del perfil.",
        ],
        "score": "20",
    },
    {
        "student": "Fabricio Jose Rivera Rupay\nU202423883\nu202423883@upc.edu.pe",
        "responsibilities": ["Apoyo al equipo en la preparación de la entrega AV1 de EnergyCore."],
        "score": "20",
    },
    {
        "student": "Renzo Zamir Revilla Quispe\nU201717085\nu201717085@upc.edu.pe",
        "responsibilities": ["Apoyo al equipo en la preparación de la entrega AV1 de EnergyCore."],
        "score": "20",
    },
    {
        "student": "Brayan Benjamin Huerta Cardenas\nU20241E550\nu20241e550@upc.edu.pe",
        "responsibilities": ["Apoyo al equipo en la preparación de la entrega AV1 de EnergyCore."],
        "score": "20",
    },
]


def set_text(cell, value, *, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell._tc.get_or_add_tcPr().append(margins)
    for edge, amount in (("top", 90), ("bottom", 90), ("left", 120), ("right", 120)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(amount))
        element.set(qn("w:type"), "dxa")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    return cell


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "666666")
        borders.append(element)


document = Document()
section = document.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(1.3)
section.right_margin = Cm(1.3)

normal = document.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(0)

rows = 4 + sum(len(member["responsibilities"]) for member in MEMBERS)
table = document.add_table(rows=rows, cols=8)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.0, 3.5, 4.6, 1.55, 1.75, 1.95, 1.55, 2.1]
for col, width in zip(table.columns, widths):
    col.width = Cm(width)
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = Cm(width)
set_borders(table)

set_text(table.cell(0, 0).merge(table.cell(0, 7)), "Participant Performance Report", bold=True, size=9)
set_text(table.cell(1, 0).merge(table.cell(1, 1)), "Nombre de Startup", bold=True)
set_text(table.cell(1, 2).merge(table.cell(1, 3)), "Teralume")
set_text(table.cell(1, 4).merge(table.cell(1, 5)), "Nombre de Producto", bold=True)
set_text(table.cell(1, 6).merge(table.cell(1, 7)), "EnergyCore")
set_text(table.cell(2, 0).merge(table.cell(2, 1)), "Entrega", bold=True)
set_text(table.cell(2, 2).merge(table.cell(2, 3)), "AV1")
set_text(table.cell(2, 4).merge(table.cell(2, 5)), "Team Leader", bold=True)
set_text(table.cell(2, 6).merge(table.cell(2, 7)), "Loa Rojas, Jean Franck")

headers = (
    "Ítem", "Estudiante", "Responsabilidades", "Cumplió a tiempo",
    "Cumplió a destiempo", "Cumplió parcialmente", "No cumplió (Cero)",
    "Calificación asignada (20 / 16 / 13 / 07 / 0)",
)
for index, heading in enumerate(headers):
    set_text(table.cell(3, index), heading, bold=True, size=6,
             align=WD_ALIGN_PARAGRAPH.CENTER)

row_index = 4
for item, member in enumerate(MEMBERS, start=1):
    responsibilities = member["responsibilities"]
    first_row = row_index
    last_row = row_index + len(responsibilities) - 1
    set_text(table.cell(first_row, 0).merge(table.cell(last_row, 0)), str(item),
             align=WD_ALIGN_PARAGRAPH.CENTER)
    set_text(table.cell(first_row, 1).merge(table.cell(last_row, 1)),
             member["student"], size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_text(table.cell(first_row, 7).merge(table.cell(last_row, 7)),
             member["score"], bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    for responsibility in responsibilities:
        set_text(table.cell(row_index, 2), responsibility)
        set_text(table.cell(row_index, 3), "X" if member["score"] else "", bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER)
        row_index += 1

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
